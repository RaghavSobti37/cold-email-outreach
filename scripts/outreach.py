#!/usr/bin/env python3
"""
cold-email-outreach — local, dedup-safe outreach tracker.

All personal data lives in a LOCAL data folder (never in this repo):
    OUTREACH_DATA=/path/to/cold-email-outreach-data   (default: ./data)

Files in the data folder
    tracker.csv        master list — one row per contact, every status change lands here
    suppression.csv    emails / domains that must never be emailed (bounced, opted out, legacy sends)
    batch_YYYY-MM-DD.json   the rendered emails picked for today's run
    runs.csv           one row per daily run (counts)
    Outreach_Tracker.xlsx   human-friendly workbook rebuilt from the CSVs

Subcommands
    init                         create the data folder + empty files
    import-history FILE.json     load previously-sent emails (from Gmail export) as do-not-contact rows
    add-leads FILE.csv|json      add new leads (deduped against tracker + suppression)
    verify                       syntax + MX + risk checks on every 'new' lead -> verified / risky / invalid
    pick --n N [--mix a=40,b=30] choose today's leads, render templates -> batch_<date>.json
    mark-sent RESULTS.json       write Gmail send results back (message/thread ids)
    mark-status STATUS.json      apply bounces / replies / auto-replies found in Gmail
    followups                    list contacts due a follow-up (no reply after N days)
    build-xlsx                   rebuild Outreach_Tracker.xlsx
    stats                        print a summary
"""
import argparse, csv, datetime as dt, json, os, re, sys, uuid
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from mxcheck import mx_ok  # noqa: E402

DATA = Path(os.environ.get("OUTREACH_DATA", ROOT / "data"))
TRACKER = DATA / "tracker.csv"
SUPP = DATA / "suppression.csv"
RUNS = DATA / "runs.csv"

FIELDS = [
    "id", "email", "contact_name", "company", "domain", "opportunity", "track", "location",
    "source_type", "source_url", "source_note", "date_found",
    "mx_ok", "published_on_source", "verify_status", "verify_note", "priority",
    "status", "template", "subject", "date_sent", "gmail_message_id", "gmail_thread_id",
    "bounce_reason", "reply_date", "reply_snippet", "followup_count", "next_action_date", "notes",
]
SUPP_FIELDS = ["value", "kind", "reason", "date_added"]
RUN_FIELDS = ["date", "picked", "sent", "bounced", "delivered", "replies_total", "notes"]

TRACKS = {
    "video_india": "Video / DOP / editing projects in India",
    "creative_tech": "Creative technologist roles & projects (anywhere)",
    "tech_remote": "Full-stack / front-end / product engineering (remote, anywhere)",
    "editing_remote": "Remote video editing / post (anywhere)",
    "adventure_video": "Travel / adventure / outdoor / documentary crew (anywhere)",
    "coding_freelance": "Freelance / white-label web dev for agencies & studios (anywhere)",
}
FINAL = {"sent", "bounced", "replied", "interested", "rejected", "no_reply", "follow_up_sent",
         "do_not_contact", "legacy_sent"}
# addresses that are almost never read by a human or bounce often
RISKY_LOCAL = re.compile(r"^(no-?reply|donotreply|do-not-reply|mailer-daemon|postmaster|abuse|privacy|legal|"
                         r"security|billing|invoice|accounts?|support|help|helpdesk|webmaster|admin|root|"
                         r"unsubscribe|bounce|noc|dmca|compliance|reasonable-accommodations|media-?enquiries|press)$")
EMAIL_RE = re.compile(r"^[a-z0-9][a-z0-9._%+\-]{0,63}@[a-z0-9\-]+(\.[a-z0-9\-]+)+$")
FREE = {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "yahoo.in", "rediffmail.com", "icloud.com",
        "live.com", "aol.com", "protonmail.com", "pm.me", "gmx.de", "yahoo.co.in"}
TYPO_DOMAINS = {"gmial.com", "gamil.com", "gmail.co", "gnail.com", "yaho.com", "hotmial.com", "outlok.com"}
DOMAIN_COOLDOWN_DAYS = 60  # don't email another address at the same company within this window


def today():
    return dt.date.today().isoformat()


def read_csv(p, fields):
    if not p.exists():
        return []
    with open(p, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        for k in fields:
            r.setdefault(k, "")
    return rows


def write_csv(p, rows, fields):
    tmp = p.with_suffix(".tmp")
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    os.replace(tmp, p)


def norm_email(e):
    e = (e or "").strip().strip(".,;:<>()[]\"'").lower()
    e = e.replace("mailto:", "")
    return e


def load():
    return read_csv(TRACKER, FIELDS), read_csv(SUPP, SUPP_FIELDS)


def blocked_sets(tracker, supp):
    emails = {r["email"] for r in tracker}
    s_emails = {r["value"] for r in supp if r["kind"] == "email"}
    s_domains = {r["value"] for r in supp if r["kind"] == "domain"}
    return emails, s_emails, s_domains


def recent_domains(tracker, days=DOMAIN_COOLDOWN_DAYS):
    cutoff = (dt.date.today() - dt.timedelta(days=days)).isoformat()
    out = set()
    for r in tracker:
        if r["status"] in FINAL and r["domain"] not in FREE and (r["date_sent"] or "0000") >= cutoff:
            out.add(r["domain"])
    return out


# ---------------------------------------------------------------- commands
def cmd_init(a):
    DATA.mkdir(parents=True, exist_ok=True)
    for p, f in ((TRACKER, FIELDS), (SUPP, SUPP_FIELDS), (RUNS, RUN_FIELDS)):
        if not p.exists():
            write_csv(p, [], f)
    print(f"data folder ready: {DATA}")


def cmd_import_history(a):
    tracker, supp = load()
    rows = json.load(open(a.file, encoding="utf-8"))
    have = {r["email"] for r in tracker}
    s_have = {(r["kind"], r["value"]) for r in supp}
    added = 0
    by_email = {}
    for r in rows:
        e = norm_email(r["email"])
        if not EMAIL_RE.match(e):
            continue
        cur = by_email.setdefault(e, {"first": r["date"], "last": r["date"], "n": 0, "bounced": False,
                                      "replied": False, "subject": r.get("subject", ""), "thread": r.get("thread_id", ""),
                                      "reply_from": r.get("reply_from", "")})
        cur["n"] += 1
        cur["first"] = min(cur["first"], r["date"]); cur["last"] = max(cur["last"], r["date"])
        cur["bounced"] |= bool(r.get("bounced")); cur["replied"] |= bool(r.get("replied"))
    for e, c in by_email.items():
        if e in have:
            continue
        status = "bounced" if c["bounced"] else ("replied" if c["replied"] else "legacy_sent")
        tracker.append({**{k: "" for k in FIELDS}, "id": uuid.uuid4().hex[:10], "email": e,
                        "domain": e.split("@")[1], "track": "legacy", "source_type": "gmail_sent_history",
                        "source_note": f"sent {c['n']}x before this system ({c['first']}..{c['last']})",
                        "date_found": c["first"], "status": status, "subject": c["subject"],
                        "date_sent": c["last"], "gmail_thread_id": c["thread"], "verify_status": "n/a",
                        "reply_snippet": c["reply_from"]})
        if ("email", e) not in s_have:
            supp.append({"value": e, "kind": "email", "reason": status, "date_added": today()})
        added += 1
    write_csv(TRACKER, tracker, FIELDS); write_csv(SUPP, supp, SUPP_FIELDS)
    print(f"imported {added} historical recipients as do-not-contact")


def _iter_leads(path):
    p = Path(path)
    if p.suffix == ".json":
        data = json.load(open(p, encoding="utf-8"))
        return data if isinstance(data, list) else data.get("leads", [])
    with open(p, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def cmd_add_leads(a):
    tracker, supp = load()
    emails, s_emails, s_domains = blocked_sets(tracker, supp)
    added = skipped = 0
    batch_domains = set()
    for L in _iter_leads(a.file):
        e = norm_email(L.get("email"))
        if not e or e in emails or e in s_emails or e.split("@")[-1] in s_domains:
            skipped += 1
            continue
        d = e.split("@")[-1]
        if d not in FREE and d in batch_domains and not a.allow_same_domain:
            skipped += 1  # one contact per company per import unless told otherwise
            continue
        batch_domains.add(d)
        row = {k: "" for k in FIELDS}
        row.update({k: (L.get(k) or "").strip() for k in FIELDS if k in L})
        row.update({"id": uuid.uuid4().hex[:10], "email": e, "domain": d, "status": "new",
                    "date_found": L.get("date_found") or today(), "followup_count": "0",
                    "track": L.get("track") or a.track or ""})
        if row["track"] not in TRACKS:
            print(f"  ! unknown track for {e}: {row['track']!r} — skipped"); skipped += 1; continue
        tracker.append(row); emails.add(e); added += 1
    write_csv(TRACKER, tracker, FIELDS)
    print(f"added {added} new leads, skipped {skipped} (duplicates / suppressed / same company)")


def cmd_verify(a):
    tracker, supp = load()
    _, s_emails, s_domains = blocked_sets(tracker, supp)
    cache = {}
    counts = {"verified": 0, "risky": 0, "invalid": 0}
    for r in tracker:
        if r["status"] != "new" and not (a.recheck and r["status"] == "queued"):
            continue
        e, d = r["email"], r["domain"]
        note = []
        status = "verified"
        local = e.split("@")[0]
        if not EMAIL_RE.match(e) or d in TYPO_DOMAINS:
            status = "invalid"; note.append("bad syntax/typo domain")
        elif e in s_emails or d in s_domains:
            status = "invalid"; note.append("suppressed")
        else:
            if d not in cache:
                cache[d] = mx_ok(d)
            ok, hosts = cache[d]
            r["mx_ok"] = "yes" if ok else "no"
            if not ok:
                status = "invalid"; note.append("domain has no mail server (MX/A)")
            elif any(h.endswith(".invalid") for h in hosts) and len(hosts) == 1:
                status = "invalid"; note.append("null MX")
            if status == "verified" and RISKY_LOCAL.match(local):
                status = "risky"; note.append(f"role address '{local}@' rarely read by humans")
            if status == "verified" and "+" in local and "hn" in local:
                note.append("plus-address from HN post")
            if status == "verified" and (r.get("published_on_source") or "").lower() not in ("yes", "y", "true", "1"):
                status = "risky"; note.append("address not seen published on the source page")
        r["verify_status"] = status
        r["verify_note"] = "; ".join(note)
        if status == "verified":
            r["status"] = "queued"
        elif status == "invalid":
            r["status"] = "do_not_contact"
        counts[status] += 1
    write_csv(TRACKER, tracker, FIELDS)
    print("verify:", counts)


# ---- templates
def load_templates():
    t = {}
    for p in (ROOT / "templates").glob("*.md"):
        raw = p.read_text(encoding="utf-8")
        m = re.match(r"---\n(.*?)\n---\n(.*)", raw, re.S)
        meta, body = {}, raw
        if m:
            for line in m.group(1).splitlines():
                if ":" in line:
                    k, v = line.split(":", 1); meta[k.strip()] = v.strip()
            body = m.group(2).strip() + "\n"
        t[p.stem] = (meta, body)
    return t


def load_profile():
    for name in ("profile.json", "profile.example.json"):
        for base in (DATA, ROOT / "config"):
            p = base / name
            if p.exists():
                return json.load(open(p, encoding="utf-8"))
    raise SystemExit("no profile.json found (copy config/profile.example.json into your data folder)")


def render(text, ctx):
    def rep(m):
        key = m.group(1).strip()
        return str(ctx.get(key, ""))
    out = re.sub(r"\{\{\s*([\w.]+)\s*\}\}", rep, text)
    return re.sub(r"\n{3,}", "\n\n", out)


LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")


def to_plain(text):
    """[label](url) -> plain text. A line made of several links (the signature row) becomes one 'label: url' per line."""
    out = []
    for line in text.split("\n"):
        links = LINK.findall(line)
        rest = LINK.sub("", line).replace("·", "").replace("|", "").strip()
        if len(links) >= 2 and not rest:
            out += [f"{lab}: {url}" for lab, url in links]
        else:
            out.append(LINK.sub(lambda m: m.group(2) if "." in m.group(1) and m.group(1).lower().rstrip("/") in m.group(2).lower()
                                else f"{m.group(1)} ({m.group(2)})", line))
    return "\n".join(out)


def to_html(text):
    """Minimal, personal-looking HTML (no images, no tracking): paragraphs, bullet lists, labelled links."""
    import html as _h
    def inline(s):
        s = _h.escape(s, quote=False)
        return LINK.sub(lambda m: f'<a href="{m.group(2)}">{m.group(1)}</a>', s)
    blocks, parts = text.strip().split("\n\n"), []
    for b in blocks:
        lines = b.split("\n")
        if lines and all(l.startswith("- ") for l in lines if l.strip()) and lines[0].strip():
            items = "".join(f"<li>{inline(l[2:])}</li>" for l in lines if l.strip())
            parts.append(f'<ul style="margin:0 0 12px 18px;padding:0">{items}</ul>')
        else:
            head, list_lines = [], []
            for l in lines:
                (list_lines if l.startswith("- ") else head).append(l)
            if head and list_lines and lines[len(head)].startswith("- "):
                parts.append(f'<p style="margin:0 0 6px">{"<br>".join(inline(x) for x in head)}</p>')
                items = "".join(f"<li>{inline(l[2:])}</li>" for l in list_lines)
                parts.append(f'<ul style="margin:0 0 12px 18px;padding:0">{items}</ul>')
            else:
                parts.append(f'<p style="margin:0 0 12px">{"<br>".join(inline(x) for x in lines)}</p>')
    return ('<div style="font-family:Arial,Helvetica,sans-serif;font-size:14px;line-height:1.5;color:#222">'
            + "".join(parts) + "</div>")


def greeting_name(r):
    n = (r.get("contact_name") or "").strip()
    if n and len(n.split()[0]) > 1 and "team" not in n.lower():
        return n.split()[0]
    c = re.sub(r"\s*\(.*?\)", "", r.get("company") or "").strip()
    return f"{c} team" if c else "there"


def cmd_pick(a):
    tracker, supp = load()
    profile = load_profile()
    templates = load_templates()
    cool = recent_domains(tracker)
    mix = {}
    if a.mix:
        for part in a.mix.split(","):
            k, v = part.split("="); mix[k.strip()] = int(v)
    else:
        mix = {k: a.n // len(TRACKS) for k in TRACKS}
    queued = [r for r in tracker if r["status"] == "queued" and r["verify_status"] == "verified"]
    queued.sort(key=lambda r: (-(int(r["priority"] or 0)), r["date_found"]))
    picked, used_domains = [], set()
    for track, n in mix.items():
        pool = [r for r in queued if r["track"] == track]
        for r in pool:
            if sum(1 for p in picked if p["track"] == track) >= n:
                break
            d = r["domain"]
            if d not in FREE and (d in cool or d in used_domains):
                continue
            picked.append(r); used_domains.add(d)
    batch = []
    for r in picked:
        tname = r.get("template") or profile["track_templates"].get(r["track"], r["track"])
        meta, body = templates[tname]
        ctx = {**profile, **{k: r[k] for k in FIELDS}, "first_name": greeting_name(r),
               "opportunity": r["opportunity"] or profile["default_opportunity"].get(r["track"], ""),
               "company": re.sub(r"\s*\(.*?\)", "", r["company"]).strip() or r["domain"].split(".")[0].title()}
        ctx["hook"] = r["notes"].split("HOOK:", 1)[1].split("|")[0].strip() if "HOOK:" in r["notes"] else ""
        ctx["hook_line"] = (ctx["hook"] + "\n\n") if ctx["hook"] else ""
        subject = render(meta.get("subject", "Quick note from {{name}}"), ctx).strip()
        raw = render(body, ctx).strip() + "\n"
        batch.append({"id": r["id"], "to": r["email"], "track": r["track"], "company": ctx["company"],
                      "subject": subject, "body": to_plain(raw), "html": to_html(raw), "template": tname})
        r["template"], r["subject"] = tname, subject
    out = DATA / f"batch_{today()}.json"
    json.dump(batch, open(out, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    write_csv(TRACKER, tracker, FIELDS)
    by = {}
    for b in batch:
        by[b["track"]] = by.get(b["track"], 0) + 1
    print(f"picked {len(batch)} -> {out}  {by}  (queued remaining: {len(queued) - len(batch)})")


def cmd_mark_sent(a):
    tracker, _ = load()
    res = json.load(open(a.file, encoding="utf-8"))
    idx = {r["id"]: r for r in tracker}
    n = 0
    for x in res:
        r = idx.get(x["id"])
        if not r:
            continue
        if x.get("ok"):
            r.update({"status": "sent", "date_sent": x.get("date") or today(),
                      "gmail_message_id": x.get("message_id", ""), "gmail_thread_id": x.get("thread_id", ""),
                      "next_action_date": (dt.date.today() + dt.timedelta(days=5)).isoformat()})
            n += 1
        else:
            r["notes"] = (r["notes"] + f" | send error {today()}: {x.get('error', '')}").strip(" |")
    write_csv(TRACKER, tracker, FIELDS)
    print(f"marked {n} sent")


def cmd_mark_status(a):
    """STATUS.json: [{"email":..., "status":"bounced|replied|interested|rejected|auto_reply|do_not_contact",
                      "reason":..., "snippet":..., "date":...}]"""
    tracker, supp = load()
    upd = json.load(open(a.file, encoding="utf-8"))
    by_email = {}
    for r in tracker:
        by_email.setdefault(r["email"], []).append(r)
    s_have = {(s["kind"], s["value"]) for s in supp}
    n = 0
    for u in upd:
        e = norm_email(u["email"])
        for r in by_email.get(e, []):
            st = u["status"]
            if st == "bounced":
                r["status"] = "bounced"; r["bounce_reason"] = u.get("reason", "")[:200]
                if ("email", e) not in s_have:
                    supp.append({"value": e, "kind": "email", "reason": "bounced", "date_added": today()})
                    s_have.add(("email", e))
            elif st == "auto_reply":
                # keep status "sent" so the day-5 follow-up still goes out; just record the acknowledgement
                if "auto-reply" not in r["notes"]:
                    r["notes"] = (r["notes"] + f" | auto-reply {u.get('date', today())}").strip(" |")
                r["reply_date"] = u.get("date", today())
                r["reply_snippet"] = ("AUTO-REPLY: " + (u.get("snippet") or u.get("reason") or ""))[:300]
            else:
                r["status"] = st
                r["reply_date"] = u.get("date", today())
                r["reply_snippet"] = (u.get("snippet") or u.get("reason") or "")[:300]
                r["next_action_date"] = ""
                if st == "do_not_contact" and ("email", e) not in s_have:
                    supp.append({"value": e, "kind": "email", "reason": "opted out", "date_added": today()})
            n += 1
    write_csv(TRACKER, tracker, FIELDS); write_csv(SUPP, supp, SUPP_FIELDS)
    print(f"updated {n} rows")


def cmd_followups(a):
    tracker, _ = load()
    due = [r for r in tracker if r["status"] == "sent" and r["next_action_date"] and r["next_action_date"] <= today()
           and int(r["followup_count"] or 0) < 1]
    profile, templates = load_profile(), load_templates()
    _, fbody = templates["followup"]
    out = []
    for r in due:
        ctx = {**profile, "first_name": greeting_name(r),
               "company": re.sub(r"\s*\(.*?\)", "", r["company"]).strip() or r["domain"].split(".")[0].title()}
        raw = render(fbody, ctx).strip() + "\n"
        out.append({"id": r["id"], "to": r["email"], "thread_id": r["gmail_thread_id"], "company": ctx["company"],
                    "track": r["track"], "first_name": ctx["first_name"],
                    "subject": "Re: " + (r["subject"] or ""), "body": to_plain(raw), "html": to_html(raw)})
    json.dump(out, open(DATA / f"followups_{today()}.json", "w"), indent=1, ensure_ascii=False)
    print(f"{len(due)} follow-ups due -> followups_{today()}.json")


def cmd_mark_followup(a):
    """FILE: [{"id":..., "ok":true, "message_id":...}] — follow-ups sent in-thread."""
    tracker, _ = load()
    idx = {r["id"]: r for r in tracker}
    n = 0
    for x in json.load(open(a.file, encoding="utf-8")):
        r = idx.get(x["id"])
        if r and x.get("ok"):
            r["followup_count"] = str(int(r["followup_count"] or 0) + 1)
            r["status"] = "follow_up_sent"
            r["next_action_date"] = (dt.date.today() + dt.timedelta(days=7)).isoformat()
            r["notes"] = (r["notes"] + f" | follow-up {today()}").strip(" |")
            n += 1
    write_csv(TRACKER, tracker, FIELDS)
    print(f"marked {n} follow-ups")


def cmd_expire(a):
    """Contacts with no reply 7+ days after their follow-up become no_reply."""
    tracker, _ = load()
    n = 0
    for r in tracker:
        if r["status"] == "follow_up_sent" and r["next_action_date"] and r["next_action_date"] <= today():
            r["status"] = "no_reply"; n += 1
    write_csv(TRACKER, tracker, FIELDS)
    print(f"{n} contacts moved to no_reply")


def cmd_sent_list(a):
    """Print sent/follow-up rows (email, thread id) so the agent can check Gmail for replies & bounces."""
    tracker, _ = load()
    rows = [{"id": r["id"], "email": r["email"], "thread_id": r["gmail_thread_id"], "date_sent": r["date_sent"],
             "status": r["status"]} for r in tracker if r["status"] in ("sent", "follow_up_sent")]
    json.dump(rows, open(DATA / "open_threads.json", "w"), indent=1)
    print(f"{len(rows)} open threads -> open_threads.json")


def cmd_log_run(a):
    runs = read_csv(RUNS, RUN_FIELDS)
    tracker, _ = load()
    runs.append({"date": today(), "picked": a.picked, "sent": a.sent, "bounced": a.bounced,
                 "delivered": int(a.sent) - int(a.bounced),
                 "replies_total": sum(1 for r in tracker if r["status"] in ("replied", "interested", "rejected")),
                 "notes": a.notes})
    write_csv(RUNS, runs, RUN_FIELDS)
    print("run logged")


def cmd_stats(a):
    tracker, supp = load()
    from collections import Counter
    print("tracker rows:", len(tracker), " suppression:", len(supp))
    print("by status:", dict(Counter(r["status"] for r in tracker)))
    print("by track (non-legacy):", dict(Counter(r["track"] for r in tracker if r["track"] != "legacy")))
    print("by source:", dict(Counter(r["source_type"] for r in tracker if r["track"] != "legacy")))


def cmd_build_xlsx(a):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.formatting.rule import FormulaRule
    tracker, supp = load()
    runs = read_csv(RUNS, RUN_FIELDS)
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1F3A5F")

    def sheet(ws, fields, rows, widths=None, table=None):
        ws.append(fields)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, Alignment(vertical="center")
        for r in rows:
            ws.append([r.get(f, "") for f in fields])
        ws.freeze_panes = "A2"
        for i, f in enumerate(fields, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(f, 16)
        if table and rows:
            t = Table(displayName=table, ref=f"A1:{get_column_letter(len(fields))}{len(rows) + 1}")
            t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)

    # Dashboard
    ws = wb.active; ws.title = "Dashboard"
    ws["A1"] = "Cold Email Outreach — Dashboard"; ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Rebuilt {dt.datetime.now():%Y-%m-%d %H:%M} from tracker.csv (live formulas below read the Contacts sheet)"
    statuses = ["queued", "sent", "bounced", "replied", "interested", "rejected", "no_reply", "follow_up_sent",
                "do_not_contact", "legacy_sent"]
    ws["A4"] = "Status"; ws["B4"] = "Contacts"
    for c in (ws["A4"], ws["B4"]):
        c.font, c.fill = head, fill
    n = len(tracker) + 1
    for i, s in enumerate(statuses, 5):
        ws.cell(i, 1, s); ws.cell(i, 2, f'=COUNTIF(Contacts!$R$2:$R${n},"{s}")')
    r0 = 5 + len(statuses) + 1
    ws.cell(r0, 1, "Track").font = head; ws.cell(r0, 1).fill = fill
    for j, s in enumerate(["queued", "sent", "bounced", "replied/interested"], 2):
        ws.cell(r0, j, s).font = head; ws.cell(r0, j).fill = fill
    for i, t in enumerate(list(TRACKS) + ["legacy"], r0 + 1):
        ws.cell(i, 1, t)
        ws.cell(i, 2, f'=COUNTIFS(Contacts!$G$2:$G${n},"{t}",Contacts!$R$2:$R${n},"queued")')
        ws.cell(i, 3, f'=COUNTIFS(Contacts!$G$2:$G${n},"{t}",Contacts!$R$2:$R${n},"sent")+COUNTIFS(Contacts!$G$2:$G${n},"{t}",Contacts!$R$2:$R${n},"follow_up_sent")+COUNTIFS(Contacts!$G$2:$G${n},"{t}",Contacts!$R$2:$R${n},"no_reply")')
        ws.cell(i, 4, f'=COUNTIFS(Contacts!$G$2:$G${n},"{t}",Contacts!$R$2:$R${n},"bounced")')
        ws.cell(i, 5, f'=COUNTIFS(Contacts!$G$2:$G${n},"{t}",Contacts!$R$2:$R${n},"replied")+COUNTIFS(Contacts!$G$2:$G${n},"{t}",Contacts!$R$2:$R${n},"interested")')
    ws.column_dimensions["A"].width = 22
    for col in "BCDE":
        ws.column_dimensions[col].width = 14

    widths = {"email": 32, "company": 24, "opportunity": 34, "source_url": 40, "source_note": 30, "subject": 40,
              "verify_note": 30, "reply_snippet": 40, "notes": 40}
    live = [r for r in tracker if r["track"] != "legacy"]
    legacy = [r for r in tracker if r["track"] == "legacy"]
    sheet(wb.create_sheet("Contacts"), FIELDS, live + legacy, widths, "Contacts")
    wsC = wb["Contacts"]
    colors = {"bounced": "F8D7DA", "replied": "D4EDDA", "interested": "C3E6CB", "sent": "E8F0FE",
              "queued": "FFF3CD", "do_not_contact": "E2E3E5"}
    for s, col in colors.items():
        wsC.conditional_formatting.add(f"A2:{get_column_letter(len(FIELDS))}{n}",
                                       FormulaRule(formula=[f'$R2="{s}"'], fill=PatternFill("solid", fgColor=col)))
    # Sources summary
    from collections import Counter, defaultdict
    src = defaultdict(Counter)
    for r in live:
        key = (r["source_type"], r["source_url"].split("/")[2] if "//" in r["source_url"] else r["source_url"])
        src[key][r["status"]] += 1
    rows = []
    for (t, host), c in sorted(src.items(), key=lambda x: -sum(x[1].values())):
        rows.append({"source_type": t, "source": host, "contacts": sum(c.values()), "sent": c["sent"],
                     "bounced": c["bounced"], "replied": c["replied"] + c["interested"], "queued": c["queued"]})
    sheet(wb.create_sheet("Sources"), ["source_type", "source", "contacts", "sent", "bounced", "replied", "queued"],
          rows, {"source": 40, "source_type": 22}, "Sources")
    sheet(wb.create_sheet("Daily Runs"), RUN_FIELDS, runs, {"notes": 60}, "Runs")
    sheet(wb.create_sheet("Do Not Contact"), SUPP_FIELDS, supp, {"value": 36, "reason": 24}, "Suppression")
    out = DATA / "Outreach_Tracker.xlsx"
    wb.save(out)
    print(f"wrote {out}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    sp.add_parser("init").set_defaults(f=cmd_init)
    p = sp.add_parser("import-history"); p.add_argument("file"); p.set_defaults(f=cmd_import_history)
    p = sp.add_parser("add-leads"); p.add_argument("file"); p.add_argument("--track")
    p.add_argument("--allow-same-domain", action="store_true"); p.set_defaults(f=cmd_add_leads)
    p = sp.add_parser("verify"); p.add_argument("--recheck", action="store_true"); p.set_defaults(f=cmd_verify)
    p = sp.add_parser("pick"); p.add_argument("--n", type=int, default=150); p.add_argument("--mix")
    p.set_defaults(f=cmd_pick)
    p = sp.add_parser("mark-sent"); p.add_argument("file"); p.set_defaults(f=cmd_mark_sent)
    p = sp.add_parser("mark-status"); p.add_argument("file"); p.set_defaults(f=cmd_mark_status)
    sp.add_parser("followups").set_defaults(f=cmd_followups)
    p = sp.add_parser("mark-followup"); p.add_argument("file"); p.set_defaults(f=cmd_mark_followup)
    sp.add_parser("expire").set_defaults(f=cmd_expire)
    sp.add_parser("open-threads").set_defaults(f=cmd_sent_list)
    p = sp.add_parser("log-run"); p.add_argument("--picked", default=0); p.add_argument("--sent", default=0)
    p.add_argument("--bounced", default=0); p.add_argument("--notes", default=""); p.set_defaults(f=cmd_log_run)
    sp.add_parser("stats").set_defaults(f=cmd_stats)
    sp.add_parser("build-xlsx").set_defaults(f=cmd_build_xlsx)
    a = ap.parse_args()
    a.f(a)


if __name__ == "__main__":
    main()
